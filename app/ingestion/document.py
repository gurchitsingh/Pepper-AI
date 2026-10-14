import os
import shutil

import pillow_heif
import pypdfium2 as pdfium
import pytesseract
from docx import Document as DocxDocument
from PIL import Image
from pypdf import PdfReader

from app.ingestion.common import IngestResult, IngestionError
from app.ingestion.transliteration import transliterate_devanagari

pillow_heif.register_heif_opener()

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".heic"}
SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx"} | IMAGE_EXTENSIONS

TESSERACT_FALLBACK_PATHS = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
)

TESSDATA_FALLBACK_DIRS = (
    r"C:\Program Files\Tesseract-OCR\tessdata",
    r"C:\Program Files (x86)\Tesseract-OCR\tessdata",
)

TESSERACT_INSTALL_HINT = (
    "OCR fallback requires Tesseract OCR, which isn't installed or isn't on PATH. Install it: "
    "Windows -- https://github.com/UB-Mannheim/tesseract/wiki or `winget install "
    "UB-Mannheim.TesseractOCR`; macOS -- `brew install tesseract`; Linux -- `apt install tesseract-ocr`."
)

HINDI_LANGPACK_INSTALL_HINT = (
    "Hindi OCR requires the Tesseract Hindi language pack (hin.traineddata), which isn't installed. "
    "Download it from https://github.com/tesseract-ocr/tessdata_fast/raw/main/hin.traineddata and place it in "
    "your tessdata folder (e.g. C:\\Program Files\\Tesseract-OCR\\tessdata); "
    "macOS/Linux -- `brew install tesseract-lang` or `apt install tesseract-ocr-hin`."
)


def _tessdata_dir() -> str | None:
    env_dir = os.environ.get("TESSDATA_PREFIX")
    if env_dir and os.path.isdir(env_dir):
        return env_dir
    tesseract_dir = os.path.dirname(pytesseract.pytesseract.tesseract_cmd)
    candidate = os.path.join(tesseract_dir, "tessdata")
    if os.path.isdir(candidate):
        return candidate
    for candidate in TESSDATA_FALLBACK_DIRS:
        if os.path.isdir(candidate):
            return candidate
    return None


def _ensure_tesseract(lang: str = "eng") -> None:
    """Resolves a usable tesseract binary for pytesseract, checking PATH first and then
    common Windows install locations (a fresh install may not be on PATH yet for an
    already-running process). Raises a clear, actionable error if none is found --
    never silently skips OCR or guesses. Also verifies every requested language's
    .traineddata file is present, so a missing Hindi pack fails clearly instead of
    silently falling back to English-only recognition.
    """
    if not shutil.which(pytesseract.pytesseract.tesseract_cmd):
        for candidate in TESSERACT_FALLBACK_PATHS:
            if os.path.exists(candidate):
                pytesseract.pytesseract.tesseract_cmd = candidate
                break
        else:
            raise IngestionError(TESSERACT_INSTALL_HINT)

    tessdata_dir = _tessdata_dir()
    for code in lang.split("+"):
        if tessdata_dir and not os.path.exists(os.path.join(tessdata_dir, f"{code}.traineddata")):
            if code == "hin":
                raise IngestionError(HINDI_LANGPACK_INSTALL_HINT)
            raise IngestionError(f"Tesseract language pack '{code}' not found in {tessdata_dir}.")


def _extract_text_file(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _extract_pdf(path: str) -> str:
    try:
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as e:
        raise IngestionError(f"This PDF is corrupt or unreadable: {e}") from e
    return "\n".join(pages)


def _ocr_pdf(path: str, lang: str = "eng+hin") -> str:
    """Fallback for PDFs with no extractable text layer (scanned documents, photos saved
    as PDF): rasterizes each page and runs Tesseract OCR on it, recognizing English and
    Hindi (Devanagari) text. Any Devanagari OCR output is transliterated to Roman script
    to match how Whisper-transcribed Hindi speech is stored (see transliteration.py).
    """
    _ensure_tesseract(lang)
    texts = []
    try:
        pdf = pdfium.PdfDocument(path)
    except Exception as e:
        raise IngestionError(f"This PDF is corrupt or unreadable: {e}") from e
    try:
        for page in pdf:
            image = page.render(scale=300 / 72).to_pil()
            page_text = pytesseract.image_to_string(image, lang=lang)
            texts.append(transliterate_devanagari(page_text))
    finally:
        pdf.close()
    return "\n".join(texts)


def _extract_docx(path: str) -> str:
    try:
        doc = DocxDocument(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    except Exception as e:
        raise IngestionError(f"This DOCX file is corrupt or unreadable: {e}") from e
    return "\n".join(paragraphs)


def _ocr_image(path: str, lang: str = "eng+hin") -> str:
    """OCR a standalone photo (as opposed to a scanned PDF page) using the same
    Tesseract setup as _ocr_pdf. A photo with no text (e.g. a landscape) is expected
    to produce empty output here -- that's not an error, see process_uploaded_document.
    """
    _ensure_tesseract(lang)
    try:
        image = Image.open(path)
        image.load()
    except Exception as e:
        raise IngestionError(f"This image is corrupt or unreadable: {e}") from e
    # pytesseract only recognizes a handful of image.format values (PNG/JPEG/etc) and
    # rejects anything else outright -- clearing it makes it fall back to PNG, which
    # works for any decoded image regardless of source format (e.g. HEIC).
    image.format = None
    return transliterate_devanagari(pytesseract.image_to_string(image, lang=lang))


def process_uploaded_document(file_path: str) -> IngestResult:
    if not os.path.exists(file_path):
        raise IngestionError(f"Uploaded file not found on disk: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    ocr_attempted = False

    if ext in (".md", ".txt"):
        text = _extract_text_file(file_path)
    elif ext == ".pdf":
        text = _extract_pdf(file_path).strip()
        if not text:
            ocr_attempted = True
            text = _ocr_pdf(file_path)
    elif ext == ".docx":
        text = _extract_docx(file_path)
    elif ext in IMAGE_EXTENSIONS:
        text = _ocr_image(file_path)
    else:
        raise IngestionError(
            f"Unsupported document file type: {ext or '(no extension)'}. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    text = text.strip()
    if not text:
        if ext in IMAGE_EXTENSIONS:
            # A photo with no readable text (e.g. a landscape) is still a valid save --
            # store it with empty text rather than failing, so it stays visible in Library.
            return IngestResult(raw_text="")
        if ext == ".pdf":
            detail = (
                "No extractable text found in this PDF, even after OCR."
                if ocr_attempted
                else "No extractable text found in this PDF."
            )
            raise IngestionError(f"{detail} It may be a blank page or an image with no readable text.")
        raise IngestionError(f"No extractable text found in this {ext} file.")

    return IngestResult(raw_text=text)
